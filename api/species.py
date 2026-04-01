"""
Vercel Python serverless function for species data ingestion and consolidation.
Parses the BABBUD LEG ESPECIES PDF to extract a unified species database.
Extensible for TXT (Ar_pa_ar.txt, FORRACAO.TXT) and XLSX (LEG PLANTIO) sources.
"""

import json
import os
import re
from http.server import BaseHTTPRequestHandler

import fitz  # PyMuPDF

# Base directory (project root)
_BASE_DIR = os.path.dirname(os.path.dirname(__file__))

# Normalized type mapping from PDF categories to internal types
TYPE_MAP = {
    "1. PALMEIRA": "palmeira",
    "2. ARV NATIVA": "arvore_nativa",
    "3. ARV EXÓTICA": "arvore_exotica",
    "3. ARV EXOTICA": "arvore_exotica",
    "4. ARV FRUTÍFERA": "arvore_frutifera",
    "4. ARV FRUTIFERA": "arvore_frutifera",
    "5. ARBUSTO": "arbusto",
    "6. FORRAÇÃO": "forracao",
    "6. FORRACAO": "forracao",
    "7. AQUÁTICA": "aquatica",
    "7. AQUATICA": "aquatica",
    "8. MATERIAL": "material",
}

# Broader grouping for compatibility with the existing map system
BROAD_TYPE_MAP = {
    "palmeira": "arvore",
    "arvore_nativa": "arvore",
    "arvore_exotica": "arvore",
    "arvore_frutifera": "arvore",
    "arbusto": "arbusto",
    "forracao": "forracao",
    "aquatica": "aquatica",
    "material": "material",
}

_species_cache = None


def parse_babbud_pdf(pdf_path):
    """Extract all species from the BABBUD LEG ESPECIES PDF."""
    species = []
    doc = fitz.open(pdf_path)

    for page_num in range(doc.page_count):
        page = doc[page_num]
        text = page.get_text()
        lines = text.strip().split("\n")

        # Skip header lines and footer
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            # Skip known headers/footers
            if (line.startswith("ESCRITÓRIO") or line == "COD" or line == "TIPO"
                    or line == "NOME CIENTIFICO" or line == "NOME POPULAR"
                    or line == "ORIGEM" or line == "INSOLAÇÃO"
                    or line.startswith("COD") or line.startswith("ÁRVORES")
                    or line.startswith("LEGENDA") or line.startswith("Página")):
                i += 1
                continue

            # Try to detect a species entry: code is typically 2-6 uppercase chars
            if re.match(r'^[A-Z][A-Z0-9]{1,5}(\s[A-Z]{1,3})?$', line):
                code = line.strip()
                # Read following lines for type, scientific name, common name, etc.
                entry = _parse_entry(lines, i, code)
                if entry:
                    species.append(entry)
                    i += entry.get("_lines_consumed", 1)
                    continue

            i += 1

    doc.close()
    return species


def _parse_entry(lines, start_idx, code):
    """Parse a species entry starting from the code line."""
    n = len(lines)
    idx = start_idx + 1
    tipo = None
    nome_cientifico = None
    nome_popular = None
    origem = None
    insolacao = None
    cod_antigo = None

    lines_consumed = 1

    # Next line should be the type
    if idx < n:
        candidate = lines[idx].strip()
        for pdf_type in TYPE_MAP:
            if candidate.startswith(pdf_type.split(".")[0] + ".") or candidate == pdf_type:
                tipo = candidate
                idx += 1
                lines_consumed += 1
                break
        else:
            # Type might be missing or merged; try matching known types
            for pdf_type in TYPE_MAP:
                if pdf_type in candidate:
                    tipo = pdf_type
                    idx += 1
                    lines_consumed += 1
                    break

    # Scientific name
    if idx < n:
        candidate = lines[idx].strip()
        if candidate and not candidate.startswith("Página") and not re.match(r'^[A-Z]{2,6}$', candidate):
            nome_cientifico = candidate
            idx += 1
            lines_consumed += 1

    # Common name
    if idx < n:
        candidate = lines[idx].strip()
        if candidate and not candidate.startswith("Página") and not re.match(r'^[A-Z]{2,6}$', candidate):
            # Check if it's a known field value (origem/insolacao)
            if candidate.lower() in ("nativa", "exótica adaptada", "exotica adaptada"):
                origem = candidate
                idx += 1
                lines_consumed += 1
            else:
                nome_popular = candidate
                idx += 1
                lines_consumed += 1

    # Origem (if not already captured)
    if origem is None and idx < n:
        candidate = lines[idx].strip()
        if candidate.lower() in ("nativa", "exótica adaptada", "exotica adaptada"):
            origem = candidate
            idx += 1
            lines_consumed += 1

    # Insolacao
    if idx < n:
        candidate = lines[idx].strip().lower()
        if candidate in ("sol", "sombra", "meia sombra", "sol/ meia sombra",
                         "meia sombra/sombra", "sol/meia sombra"):
            insolacao = lines[idx].strip()
            idx += 1
            lines_consumed += 1

    # Cod antigo
    if idx < n:
        candidate = lines[idx].strip()
        if re.match(r'^[A-Z][A-Z0-9]{1,5}(\s[A-Z]{1,3})?$', candidate):
            cod_antigo = candidate
            idx += 1
            lines_consumed += 1

    if not nome_cientifico:
        return None

    # Normalize type
    normalized_type = None
    broad_type = None
    if tipo:
        for pdf_type, norm in TYPE_MAP.items():
            if tipo.startswith(pdf_type.split(".")[0] + ".") or tipo == pdf_type:
                normalized_type = norm
                broad_type = BROAD_TYPE_MAP.get(norm, norm)
                break

    return {
        "code": code,
        "type": normalized_type or "desconhecido",
        "broad_type": broad_type or "desconhecido",
        "type_label": tipo or "",
        "scientific_name": nome_cientifico,
        "common_name": nome_popular or "",
        "origin": origem or "",
        "sun_exposure": insolacao or "",
        "old_code": cod_antigo or "",
        "image_url": "",  # Placeholder for XLSX image association
        "source": "pdf",
        "_lines_consumed": lines_consumed,
    }


def parse_txt_ar_pa_ar(filepath):
    """Parse Ar_pa_ar.txt structured species list. (Extension point)"""
    if not os.path.exists(filepath):
        return []
    species = []
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    # TXT format: tab or pipe delimited, fields: code, type, scientific, common, dimensions, obs, colors
    for line in content.strip().split("\n"):
        parts = re.split(r'\t+|\|', line.strip())
        if len(parts) < 4:
            continue
        code = parts[0].strip()
        if not re.match(r'^[A-Z][A-Z0-9]{1,5}$', code):
            continue
        species.append({
            "code": code,
            "type": parts[1].strip() if len(parts) > 1 else "",
            "broad_type": "",
            "type_label": parts[1].strip() if len(parts) > 1 else "",
            "scientific_name": parts[2].strip() if len(parts) > 2 else "",
            "common_name": parts[3].strip() if len(parts) > 3 else "",
            "origin": "",
            "sun_exposure": "",
            "old_code": "",
            "image_url": "",
            "source": "txt_ar_pa_ar",
        })
    return species


def parse_txt_forracao(filepath):
    """Parse FORRACAO.TXT structured list.
    Format: ("CODE" "TYPE" "Scientific name" "Common name" "H" "D" "Obs" (colors) "price" "STATUS")
    """
    if not os.path.exists(filepath):
        return []
    species = []
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    for line in content.strip().split("\n"):
        line = line.strip()
        if not line.startswith("("):
            continue
        # Extract quoted fields
        fields = re.findall(r'"([^"]*)"', line)
        if len(fields) < 4:
            continue
        code = fields[0].strip()
        if not code or not re.match(r'^[A-Z][A-Z0-9]{1,5}$', code):
            continue
        cat = fields[1].strip()  # FO, AR, PA, etc.
        scientific = fields[2].strip()
        common = fields[3].strip()
        # Height and diameter if present
        height = fields[4].strip() if len(fields) > 4 else ""
        diameter = fields[5].strip() if len(fields) > 5 else ""
        obs = fields[6].strip() if len(fields) > 6 else ""

        # Skip non-plant materials
        if common.lower() in ("areia", "casca de árvore", "casca de arvore",
                               "casca de árvore mini", "casca de arvore mini",
                               "pedra", "seixo"):
            continue

        species.append({
            "code": code,
            "type": "forracao",
            "broad_type": "forracao",
            "type_label": "6. FORRAÇÃO",
            "scientific_name": scientific,
            "common_name": common,
            "origin": "",
            "sun_exposure": "",
            "old_code": "",
            "image_url": "",
            "height": height,
            "diameter": diameter,
            "planting_notes": obs,
            "source": "txt_forracao",
        })
    return species


def load_xlsx_images(base_dir):
    """Load species -> image mapping from extracted XLSX images.
    Reads the mapping.json created during image extraction."""
    mapping_path = os.path.join(base_dir, "especies_img", "mapping.json")
    if not os.path.exists(mapping_path):
        return {}
    import json
    with open(mapping_path, "r") as f:
        return json.load(f)


def load_xlsx_species(xlsx_path):
    """Parse 'uso interno PLANILHA' sheet from XLSX for additional species data."""
    if not os.path.exists(xlsx_path):
        return []
    try:
        import openpyxl
    except ImportError:
        return []

    species = []
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if 'uso interno PLANILHA' not in wb.sheetnames:
        wb.close()
        return []

    ws = wb['uso interno PLANILHA']
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        if not re.match(r'^[A-Z][A-Z0-9]{1,5}(\s[A-Z]{1,3})?$', code):
            continue
        tipo = str(row[1]).strip() if row[1] else ""
        scientific = str(row[2]).strip() if row[2] else ""
        common = str(row[3]).strip() if row[3] else ""
        origin = str(row[6]).strip() if len(row) > 6 and row[6] else ""
        sun = str(row[7]).strip() if len(row) > 7 and row[7] else ""

        if not scientific:
            continue

        # Determine normalized type
        normalized_type = "desconhecido"
        broad_type = "desconhecido"
        for pdf_type, norm in TYPE_MAP.items():
            if tipo.startswith(pdf_type.split(".")[0] + ".") or tipo == pdf_type:
                normalized_type = norm
                broad_type = BROAD_TYPE_MAP.get(norm, norm)
                break

        species.append({
            "code": code,
            "type": normalized_type,
            "broad_type": broad_type,
            "type_label": tipo,
            "scientific_name": scientific,
            "common_name": common,
            "origin": origin,
            "sun_exposure": sun,
            "old_code": "",
            "image_url": "",
            "source": "xlsx",
        })

    wb.close()
    return species


def consolidate_species(pdf_species, txt_ar=None, txt_forr=None, xlsx_species=None, xlsx_images=None):
    """
    Consolidate species from all sources into a unified dictionary.
    Priority: PDF > XLSX > TXT (for enrichment).
    """
    unified = {}

    # 1. PDF as primary source
    for sp in pdf_species:
        code = sp["code"]
        entry = {k: v for k, v in sp.items() if not k.startswith("_")}
        unified[code] = entry

    # 2. XLSX species enrichment (second priority)
    for sp in (xlsx_species or []):
        code = sp["code"]
        if code in unified:
            for field in ("common_name", "sun_exposure", "origin"):
                if not unified[code].get(field) and sp.get(field):
                    unified[code][field] = sp[field]
            if "xlsx" not in unified[code].get("source", ""):
                unified[code]["source"] += "+xlsx"
        else:
            entry = {k: v for k, v in sp.items() if not k.startswith("_")}
            unified[code] = entry

    # 3. TXT enrichment (planting notes, dimensions)
    for txt_list in (txt_ar or [], txt_forr or []):
        for sp in txt_list:
            code = sp["code"]
            if code in unified:
                for field in ("common_name", "sun_exposure", "origin",
                              "height", "diameter", "planting_notes"):
                    if not unified[code].get(field) and sp.get(field):
                        unified[code][field] = sp[field]
                if "txt" not in unified[code].get("source", ""):
                    unified[code]["source"] += "+txt"
            else:
                entry = {k: v for k, v in sp.items() if not k.startswith("_")}
                unified[code] = entry

    # 4. XLSX image association
    if xlsx_images:
        for code, img_file in xlsx_images.items():
            upper_code = code.upper()
            if upper_code in unified:
                unified[upper_code]["image_url"] = "/especies_img/" + img_file
                if "img" not in unified[upper_code].get("source", ""):
                    unified[upper_code]["source"] += "+img"

    # 5. Filter out non-plant materials
    non_plant_codes = set()
    for code, sp in unified.items():
        if sp.get("broad_type") == "material" or sp.get("type") == "material":
            non_plant_codes.add(code)
        name = (sp.get("scientific_name", "") + sp.get("common_name", "")).lower()
        if any(m in name for m in ("areia", "casca de árvore", "casca de arvore",
                                    "pedra", "seixo", "brita", "cascalho")):
            non_plant_codes.add(code)

    for code in non_plant_codes:
        unified.pop(code, None)

    return unified


def get_species_database():
    """Load and cache the consolidated species database."""
    global _species_cache
    if _species_cache is not None:
        return _species_cache

    base_dir = os.path.dirname(os.path.dirname(__file__))

    # 1. PDF primary source
    pdf_species = []
    pdf_path = os.path.join(base_dir, "arquivos", "BABBUD LEG ESPECIES 03.2026-R00.pdf")
    if not os.path.exists(pdf_path):
        # Also check root
        pdf_path = os.path.join(base_dir, "BABBUD LEG ESPECIES 03.2026-R00.pdf")
    if os.path.exists(pdf_path):
        pdf_species = parse_babbud_pdf(pdf_path)

    # 2. TXT sources
    txt_ar = parse_txt_ar_pa_ar(os.path.join(base_dir, "Ar_pa_ar.txt"))
    if not txt_ar:
        txt_ar = parse_txt_ar_pa_ar(os.path.join(base_dir, "arquivos", "Ar_pa_ar.txt"))
    txt_forr = parse_txt_forracao(os.path.join(base_dir, "FORRACAO.TXT"))
    if not txt_forr:
        txt_forr = parse_txt_forracao(os.path.join(base_dir, "arquivos", "FORRACAO.TXT"))

    # 3. XLSX species data
    xlsx_path = None
    for candidate in [
        os.path.join(base_dir, "XXXX-XXX-PA-AP-500-R00_2026 R0A LEG PLANTIO.xlsx"),
        os.path.join(base_dir, "arquivos", "XXXX-XXX-PA-AP-500-R00_2026 R0A LEG PLANTIO.xlsx"),
    ]:
        if os.path.exists(candidate):
            xlsx_path = candidate
            break
    xlsx_species = load_xlsx_species(xlsx_path) if xlsx_path else []

    # 4. XLSX images
    xlsx_images = load_xlsx_images(base_dir)

    _species_cache = consolidate_species(
        pdf_species, txt_ar, txt_forr, xlsx_species, xlsx_images
    )
    return _species_cache


class handler(BaseHTTPRequestHandler):
    def do_GET(self):
        try:
            path = self.path.split("?")[0]
            query = self.path.split("?")[1] if "?" in self.path else ""
            params = dict(p.split("=", 1) for p in query.split("&") if "=" in p)

            db = get_species_database()

            # GET /api/species?code=PLRU — single species lookup
            if "code" in params:
                code = params["code"].upper()
                if code in db:
                    self._json_response(200, db[code])
                else:
                    self._error(404, f"Species '{code}' not found")
                return

            # GET /api/species?type=arvore — filter by broad type
            if "type" in params:
                type_filter = params["type"].lower()
                filtered = {k: v for k, v in db.items()
                           if v.get("broad_type") == type_filter or v.get("type") == type_filter}
                self._json_response(200, {
                    "count": len(filtered),
                    "type": type_filter,
                    "species": filtered,
                })
                return

            # GET /api/species?search=palm — search by name
            if "search" in params:
                q = params["search"].lower()
                filtered = {k: v for k, v in db.items()
                           if q in v.get("scientific_name", "").lower()
                           or q in v.get("common_name", "").lower()
                           or q in k.lower()}
                self._json_response(200, {
                    "count": len(filtered),
                    "query": params["search"],
                    "species": filtered,
                })
                return

            # GET /api/species — full database
            # Group by type for overview
            type_counts = {}
            for sp in db.values():
                bt = sp.get("broad_type", "desconhecido")
                type_counts[bt] = type_counts.get(bt, 0) + 1

            self._json_response(200, {
                "total": len(db),
                "types": type_counts,
                "species": db,
            })

        except Exception as e:
            self._error(500, str(e))

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors_headers()
        self.send_header("Content-Length", "0")
        self.end_headers()

    def _cors_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def _json_response(self, status, data):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self._cors_headers()
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _error(self, status, message):
        self._json_response(status, {"error": message})
